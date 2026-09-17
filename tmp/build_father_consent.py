from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE


OUT = "output/documents/Zustimmungserklaerung_Vater_NETOVANNYY_ANDREY.docx"


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_bottom_border(cell, color="B8B8B8", size="8"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "0")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def set_repeat_table_layout(table, widths):
    table.autofit = False
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = widths[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_run(p, text, bold=False, italic=False, size=None, color="000000"):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.name = "Arial"
    r._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    r._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
    r.font.color.rgb = RGBColor.from_string(color)
    if size:
        r.font.size = Pt(size)
    return r


def add_field_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_repeat_table_layout(table, [Inches(2.65), Inches(4.45)])
    for label, value in rows:
        cells = table.add_row().cells
        p0 = cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        add_run(p0, label, bold=True, size=9.6)
        p1 = cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        add_run(p1, value, size=10.1)
        set_bottom_border(cells[0])
        set_bottom_border(cells[1])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_section_label(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    add_run(p, text, bold=True, size=10.3)
    return p


def add_pair(doc, ru, de):
    p1 = doc.add_paragraph(style="Legal Body")
    p1.paragraph_format.keep_with_next = True
    add_run(p1, ru, size=10.2)
    p2 = doc.add_paragraph(style="Legal Body")
    p2.paragraph_format.space_after = Pt(7)
    add_run(p2, de, italic=True, size=10.2, color="303030")


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.55)
sec.bottom_margin = Inches(0.55)
sec.left_margin = Inches(0.7)
sec.right_margin = Inches(0.7)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
normal.font.size = Pt(10.2)
normal.font.color.rgb = RGBColor(0, 0, 0)

title_style = styles["Title"]
title_style.font.name = "Arial"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
title_style.font.size = Pt(14.5)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style.paragraph_format.space_after = Pt(4)
title_ppr = title_style._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

legal = styles.add_style("Legal Body", WD_STYLE_TYPE.PARAGRAPH)
legal.base_style = styles["Normal"]
legal.paragraph_format.space_before = Pt(0)
legal.paragraph_format.space_after = Pt(2)
legal.paragraph_format.line_spacing = 1.04
legal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
legal.paragraph_format.keep_together = True

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_run(p, "СОГЛАСИЕ ОТЦА", bold=True, size=14.5)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(9)
add_run(p, "ZUSTIMMUNGSERKLÄRUNG DES VATERS", bold=True, size=13.2)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(5)
add_run(p, "Я, нижеподписавшийся / Ich, der Unterzeichnende:", size=10.2)

add_field_table(doc, [
    ("Фамилия, имя\nVor- und Nachname", "ANDREY NETOVANNYY"),
    ("Дата рождения\nGeburtsdatum", "17.08.1986"),
    ("Гражданство\nStaatsangehörigkeit", "Российская Федерация / Russische Föderation"),
    ("Паспорт №\nReisepassnummer", "51 7277938"),
    ("Адрес проживания\nWohnadresse", "____________________________________________"),
])

add_section_label(doc, "являюсь отцом несовершеннолетнего ребёнка / bin der Vater des minderjährigen Kindes:")
add_field_table(doc, [
    ("Фамилия, имя ребёнка\nVor- und Nachname des Kindes", "IVAN NETOVANNYI"),
    ("Дата рождения\nGeburtsdatum", "27.11.2015"),
    ("Гражданство\nStaatsangehörigkeit", "Украина / Ukraine"),
    ("Паспорт №\nReisepassnummer", "GF493318"),
])

p = doc.add_paragraph(style="Legal Body")
p.paragraph_format.space_after = Pt(6)
add_run(p, "Мать ребёнка / Mutter des Kindes: ", bold=True, size=10.2)
add_run(p, "VALERIIA PECHENA, дата рождения / Geburtsdatum: 26.10.1987, гражданство / Staatsangehörigkeit: Украина / Ukraine, паспорт № / Reisepassnummer: GK241226.", size=10.2)

add_pair(
    doc,
    "Мать ребёнка и я состоим в разводе.",
    "Die Mutter des Kindes und ich sind geschieden.",
)

add_pair(
    doc,
    "Настоящим я выражаю своё добровольное и безусловное согласие на то, чтобы мой несовершеннолетний сын IVAN NETOVANNYI проживал совместно со своей матерью VALERIIA PECHENA и имел место жительства и обычного пребывания в Австрии, в том числе в федеральной земле Верхняя Австрия (Oberösterreich).",
    "Hiermit erkläre ich freiwillig und uneingeschränkt meine Zustimmung dazu, dass mein minderjähriger Sohn IVAN NETOVANNYI gemeinsam mit seiner Mutter VALERIIA PECHENA in Österreich, insbesondere im Bundesland Oberösterreich, wohnt und dort seinen Wohnsitz sowie seinen gewöhnlichen Aufenthalt begründet.",
)

add_pair(
    doc,
    "Я также не возражаю против подачи заявления, оформления и получения для моего сына разрешения на проживание в Австрии, в частности Rot-Weiß-Rot-Karte plus либо иного применимого вида на жительство, а также против проведения всех связанных с этим административных процедур.",
    "Ich erhebe außerdem keine Einwände gegen die Beantragung, Erteilung und Entgegennahme eines österreichischen Aufenthaltstitels für meinen Sohn, insbesondere einer Rot-Weiß-Rot-Karte plus oder eines anderen einschlägigen Aufenthaltstitels, sowie gegen die Durchführung aller damit verbundenen behördlichen Verfahren.",
)

add_pair(
    doc,
    "Я согласен с тем, чтобы мать ребёнка в связи с указанными процедурами представляла его интересы перед компетентными органами и учреждениями, подавала заявления и документы и получала соответствующие решения и документы в пределах, допускаемых законом.",
    "Ich bin damit einverstanden, dass die Mutter des Kindes ihn im Zusammenhang mit diesen Verfahren vor den zuständigen Behörden und Stellen vertritt, Anträge und Unterlagen einreicht sowie die entsprechenden Entscheidungen und Dokumente entgegennimmt, soweit dies gesetzlich zulässig ist.",
)

add_pair(
    doc,
    "Я подтверждаю, что мне известно о месте проживания ребёнка в Австрии и что я не возражаю против его проживания вместе с матерью в Верхней Австрии.",
    "Ich bestätige, dass mir der Wohnort meines Kindes in Österreich bekannt ist und dass ich keine Einwände dagegen habe, dass mein Kind gemeinsam mit seiner Mutter in Oberösterreich lebt.",
)

add_pair(
    doc,
    "Настоящее согласие предоставляется мной добровольно и предназначено для представления в компетентные органы Австрии.",
    "Diese Zustimmung erteile ich freiwillig; sie ist zur Vorlage bei den zuständigen österreichischen Behörden bestimmt.",
)

sig = doc.add_table(rows=3, cols=2)
sig.alignment = WD_TABLE_ALIGNMENT.CENTER
set_repeat_table_layout(sig, [Inches(3.55), Inches(3.55)])
for j, label in enumerate(("Место / Ort", "Дата / Datum")):
    c = sig.rows[0].cells[j]
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_run(p, label, bold=True, size=9.5)
    q = c.add_paragraph()
    q.paragraph_format.space_before = Pt(7)
    q.paragraph_format.space_after = Pt(0)
    add_run(q, "__________________________________", size=10.1, color="777777")

name_cell = sig.rows[1].cells[0].merge(sig.rows[1].cells[1])
p = name_cell.paragraphs[0]
p.paragraph_format.space_after = Pt(2)
add_run(p, "Фамилия, имя отца / Vor- und Nachname des Vaters", bold=True, size=9.5)
q = name_cell.add_paragraph()
q.paragraph_format.space_after = Pt(0)
add_run(q, "ANDREY NETOVANNYY", size=10.1)

sign_cell = sig.rows[2].cells[0].merge(sig.rows[2].cells[1])
p = sign_cell.paragraphs[0]
p.paragraph_format.space_after = Pt(0)
add_run(p, "Подпись отца / Unterschrift des Vaters", bold=True, size=9.5)
q = sign_cell.add_paragraph()
q.paragraph_format.space_before = Pt(13)
q.paragraph_format.space_after = Pt(0)
add_run(q, "____________________________________________", size=10.1, color="777777")

doc.add_paragraph().paragraph_format.space_after = Pt(0)
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(7)
p.paragraph_format.space_after = Pt(4)
p.paragraph_format.keep_with_next = True
add_run(p, "Нотариальное заверение подписи / Notarielle Beglaubigung der Unterschrift", bold=True, size=10.2)
for _ in range(4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    add_run(p, "________________________________________________________________________________", size=9.5, color="B0B0B0")

doc.core_properties.title = "Согласие отца на проживание ребёнка в Австрии"
doc.core_properties.subject = "Zustimmungserklärung des Vaters"
doc.core_properties.author = "ANDREY NETOVANNYY"
doc.core_properties.keywords = "Austria; Oberösterreich; Zustimmung; minderjähriges Kind"

doc.save(OUT)
print(OUT)
